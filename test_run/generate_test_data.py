"""
Generates test_sheets.pdf (4 filled OMR pages) and lista_alumnos.xlsx
in the same folder as this script.
Run from anywhere: python test_run/generate_test_data.py
"""
import os, sys

# Put omr_app on the path so we can import layout + sheet_generator
ROOT     = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP_DIR  = os.path.join(ROOT, "omr_app")
OUT_DIR  = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, APP_DIR)

from reportlab.pdfgen import canvas
from reportlab.lib import colors
import layout as L
from sheet_generator import _draw_page, _register_arial, FONT, FONT_BOLD

if _register_arial():
    FONT      = "Arial"
    FONT_BOLD = "Arial-Bold"

NUM_MC = 10   # questions in Section 1

# ── Student definitions ────────────────────────────────────────────────────────
# folio: 2-char string of digits 1-9 each (no zero)
# mc: list of 0-based choice indices (A=0 B=1 C=2 D=3), length NUM_MC
# sk: list of 0-based scale indices  (1=0 2=1 3=2 4=3 5=4), length 10
STUDENTS = [
    {
        "nombre": "Ana García López",
        "folio":  "12",           # N1=1, N2=2
        "grado":  "2",            # index 1
        "grupo":  "B",            # index 1
        "mc":     [0,1,2,3,1,0,1,2,3,0],   # A B C D B A B C D A
        "sk":     [3,3,2,4,3,2,3,4,3,2],   # 4 4 3 5 4 3 4 5 4 3
        # words: ADN, HERENCIA, PROTEÍNAS, GENES
        "words":  [0, 1, 4, 5],
    },
    {
        "nombre": "Carlos Mendoza Ruiz",
        "folio":  "35",           # N1=3, N2=5
        "grado":  "2",
        "grupo":  "B",
        "mc":     [0,1,2,3,1,0,1,2,3,1],   # A B C D B A B C D B (9/10)
        "sk":     [2,2,3,3,2,2,3,2,2,3],   # 3 3 4 4 3 3 4 3 3 4
        # words: ADN, PATERNIDAD, MICROPIPETA, PRECISO
        "words":  [0, 2, 6, 10],
    },
    {
        "nombre": "Sofia Torres Vega",
        "folio":  "47",           # N1=4, N2=7
        "grado":  "2",
        "grupo":  "B",
        "mc":     [0,1,0,3,1,0,1,2,3,0],   # A B A D B A B C D A (9/10, Q3 wrong)
        "sk":     [4,4,3,4,4,4,3,4,4,3],   # 5 5 4 5 5 5 4 5 5 4
        # words: ADN, HERENCIA, PROTEÍNAS, CAMPO ELÉCTRICO, MOLÉCULAS
        "words":  [0, 1, 4, 7, 8],
    },
    {
        "nombre": "Roberto Jiménez Cruz",
        "folio":  "63",           # N1=6, N2=3
        "grado":  "2",
        "grupo":  "B",
        "mc":     [1,1,2,3,1,0,1,2,3,0],   # B B C D B A B C D A (9/10, Q1 wrong)
        "sk":     [1,2,2,1,2,2,1,2,2,1],   # 2 3 3 2 3 3 2 3 3 2
        # words: ADN, SEPARACIÓN, GENES, MEMBRANA, MICROLITROS
        "words":  [0, 3, 5, 9, 11],
    },
]


def _fill_bubble(c, x, y, r=L.BUBBLE_R):
    """Draw a filled black bubble."""
    c.setFillColor(colors.black)
    c.setStrokeColor(colors.black)
    c.circle(x, y, r, stroke=0, fill=1)
    c.setFillColor(colors.black)


def _draw_filled_page(c, student, exam_name):
    # 1. Draw the blank template
    _draw_page(c, NUM_MC, exam_name)

    f  = student["folio"]   # e.g. "35"
    d1 = int(f[0])           # 0-based index into FOLIO_VALUES (0→idx0, 1→idx1…)
    d2 = int(f[1])

    # 2. Grado bubble
    g_idx = L.GRADO_VALUES.index(student["grado"])
    _fill_bubble(c, L.grado_bubble_x(g_idx), L.BUBBLE_ROW_Y)

    # 3. Grupo bubble
    grp_idx = L.GRUPO_VALUES.index(student["grupo"])
    _fill_bubble(c, L.grupo_bubble_x(grp_idx), L.BUBBLE_ROW_Y)

    # 4. Folio — N1 (row 1) and N2 (row 2)
    _fill_bubble(c, L.folio_bubble_x(d1), L.FOLIO_ROW1_Y)
    _fill_bubble(c, L.folio_bubble_x(d2), L.FOLIO_ROW2_Y)

    # 5. Section 1 — MC answers
    for q_idx, choice_idx in enumerate(student["mc"]):
        bx, by = L.sec1_bubble(q_idx, choice_idx, NUM_MC)
        _fill_bubble(c, bx, by)

    # 6. Section 2 — SK answers
    for q_idx, choice_idx in enumerate(student["sk"]):
        bx, by = L.sec2_bubble(q_idx, choice_idx, NUM_MC)
        _fill_bubble(c, bx, by)

    # 7. Section 3 — Word selections
    for word_idx in student.get("words", []):
        bx, by = L.sec3_bubble(word_idx, NUM_MC)
        _fill_bubble(c, bx, by)


def generate_pdf(out_path, exam_name="Examen Test"):
    c = canvas.Canvas(out_path, pagesize=(L.PAGE_W, L.PAGE_H))
    for i, student in enumerate(STUDENTS):
        _draw_filled_page(c, student, exam_name)
        if i < len(STUDENTS) - 1:
            c.showPage()
    c.save()
    print(f"PDF generado: {out_path}")


def generate_student_list(out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Alumnos"

    # Header
    ws["A1"] = "Folio"
    ws["B1"] = "Nombre"
    for cell in (ws["A1"], ws["B1"]):
        cell.font      = Font(bold=True, name="Arial")
        cell.fill      = PatternFill("solid", fgColor="2C3E50")
        cell.font      = Font(bold=True, color="FFFFFF", name="Arial")
        cell.alignment = Alignment(horizontal="center")

    for student in STUDENTS:
        ws.append([student["folio"], student["nombre"]])

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 30
    wb.save(out_path)
    print(f"Lista guardada: {out_path}")


if __name__ == "__main__":
    generate_pdf(os.path.join(OUT_DIR, "test_sheets.pdf"))
    generate_student_list(os.path.join(OUT_DIR, "lista_alumnos.xlsx"))
    print("\nListo. Abre el app y:")
    print("  PDF:         test_run/test_sheets.pdf")
    print("  Alumnos:     test_run/lista_alumnos.xlsx")
    print(f"  Preguntas:   {NUM_MC}")
    print("  Clave MC:    A B C D B A B C D A")
    print("  Grado: 2 | Grupo: B | Folios: 12, 35, 47, 63")
